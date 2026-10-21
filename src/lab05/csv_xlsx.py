import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    source = Path(csv_path)
    target = Path(xlsx_path)

    if not source.is_file():
        raise FileNotFoundError("Файл не найден")
    if source.suffix.lower() != ".csv":
        raise ValueError("Входной файл должен быть CSV")
    if target.suffix.lower() != ".xlsx":
        raise ValueError("Выходной файл должен быть XLSX")
    
    rows = []
    with source.open(encoding="utf-8", newline="") as f:
        for row in csv.reader(f):
            if row:
                rows.append(row)

    if len(rows) < 2:
        raise ValueError("В CSV нет данных")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in rows:
        ws.append(row)

    widths = {}
    for row in rows:
        for i, value in enumerate(row, start=1):
            widths[i] = max(widths.get(i, 0), len(value))

    for i, width in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = max(width, 8)

    wb.save(xlsx_path)
